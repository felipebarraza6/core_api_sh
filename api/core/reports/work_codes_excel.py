"""
Generador de Excel para el reporte de códigos de obra.
"""

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter


# Colores
HEADER_FILL = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
HEADER_FONT = Font(color="FFFFFF", bold=True)
SUPERADO_FILL = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")
SUPERADO_FONT = Font(color="9C0006")
PROXIMO_FILL = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")
PROXIMO_FONT = Font(color="9C5700")
SIN_CONFIG_FILL = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")


def generate_work_codes_excel(rows, year, output_path=None):
    """
    Genera un archivo Excel con el reporte de códigos de obra.

    Args:
        rows: Lista de diccionarios devuelta por build_work_codes_report.
        year: Año del reporte (int).
        output_path: Ruta donde guardar el archivo. Si es None, retorna el Workbook.

    Returns:
        str: Ruta del archivo guardado, o Workbook si output_path es None.
    """
    wb = Workbook()
    ws = wb.active
    ws.title = f"CodigosObra_{year}"

    headers = [
        "Cliente",
        "Proyecto",
        "Punto ID",
        "Punto",
        "Código Obra",
        "Estándar",
        "Caudal Aut. (lt/s)",
        "Veces Sup. Caudal",
        "Caudal Prom. Año (lt/s)",
        "Total Aut. (m3)",
        "Consumo Año (m3)",
        "% Gastado",
        "% Disponible",
        "m3 Disponibles",
        "Estado Total",
        "Estado Caudal",
        "Última Medición Año",
        "Última Medición Punto",
    ]

    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

    for r in rows:
        ws.append([
            r["cliente"],
            r["proyecto"],
            r["punto_id"],
            r["punto"],
            r["codigo_obra"],
            r["estandar"],
            r["caudal_autorizado_lt_s"],
            r["veces_supero_caudal"],
            r["caudal_promedio_lt_s"],
            r["total_autorizado_m3"],
            r["consumo_acumulado_m3"],
            r["pct_gastado"],
            r["pct_disponible"],
            r["m3_disponibles"],
            r["estado_total"],
            r["estado_caudal"],
            r["ultima_medicion_anio"],
            r["ultima_medicion_punto"],
        ])

    # Formato condicional simple por fila
    for idx, row in enumerate(ws.iter_rows(min_row=2, max_row=ws.max_row), start=2):
        estado_total = rows[idx - 2]["estado_total"]
        estado_caudal = rows[idx - 2]["estado_caudal"]

        if estado_total == "SUPERADO" or estado_caudal == "SUPERADO":
            for cell in row:
                cell.fill = SUPERADO_FILL
                cell.font = SUPERADO_FONT
        elif estado_total == "PROXIMO":
            for cell in row:
                cell.fill = PROXIMO_FILL
                cell.font = PROXIMO_FONT
        elif estado_total == "SIN_TOTAL" or estado_caudal == "SIN_CAUDAL":
            for cell in row:
                cell.fill = SIN_CONFIG_FILL

    # Ajustar anchos
    for col_idx, header in enumerate(headers, start=1):
        max_length = len(header)
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, min_col=col_idx, max_col=col_idx):
            for cell in row:
                if cell.value is not None:
                    max_length = max(max_length, len(str(cell.value)))
        ws.column_dimensions[get_column_letter(col_idx)].width = min(max_length + 2, 50)

    # Congelar primera fila
    ws.freeze_panes = "A2"

    # Filtros
    ws.auto_filter.ref = ws.dimensions

    if output_path:
        wb.save(output_path)
        return output_path

    return wb
